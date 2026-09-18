# -*- coding: utf-8 -*-
"""稟議パックのExcel 2点。数字は src/lib/pricing/facts.json（正本）から読む。

    python3 scripts/approval-pack/build_xlsx.py <リポジトリのルート>

ここに数字を直接書かない。直すのは facts.json。
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(sys.argv[1] if len(sys.argv) > 1 else '.')
F = json.loads((ROOT / 'src/lib/pricing/facts.json').read_text(encoding='utf-8'))
OUT = ROOT / 'public/docs'
OUT.mkdir(parents=True, exist_ok=True)

BLUE = Font(name='Arial', color='0000FF')
BLACK = Font(name='Arial')
BOLD = Font(name='Arial', bold=True)
H = Font(name='Arial', bold=True, size=14)
H2 = Font(name='Arial', bold=True, size=11)
NOTE = Font(name='Arial', size=9, color='808080')
WHITE_B = Font(name='Arial', bold=True, color='FFFFFF')
YELLOW = PatternFill('solid', fgColor='FFFF00')
NAVY = PatternFill('solid', fgColor='334155')
GRAY = PatternFill('solid', fgColor='F2F2F2')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical='top', wrap_text=True)
YEN = '¥#,##0;(¥#,##0);-'
HRS = '0.0'


def build_roi() -> None:
    tco = F['tco']
    pro = F['agentpm']['pro']
    wb = Workbook()
    ws = wb.active
    ws.title = '試算'
    ws.sheet_view.showGridLines = False
    for col, w in zip('ABCDE', [30, 13, 13, 13, 50]):
        ws.column_dimensions[col].width = w

    ws['A1'] = 'AgentPM 費用の試算シート'
    ws['A1'].font = H
    ws['A2'] = '黄色い枠の数字（青字）だけ書き換えてください。ほかは自動で計算されます。'
    ws['A2'].font = NOTE

    ws['A4'] = '1. 前提を入れる'
    ws['A4'].font = H2
    for r, label, val, unit, note in [
        (5, '平均の時給', tco['hourlyYen'], '円', '給与＋賞与＋社会保険を時間で割った額'),
        (6, '同時に動いている案件数', tco['projects'], '件', '相手先とやり取りしている案件の数'),
    ]:
        ws[f'A{r}'] = label
        ws[f'A{r}'].font = BLACK
        c = ws[f'B{r}']
        c.value, c.font, c.fill, c.border, c.number_format = val, BLUE, YELLOW, BOX, '#,##0'
        ws[f'C{r}'] = unit
        ws[f'C{r}'].font = NOTE
        ws[f'E{r}'] = note
        ws[f'E{r}'].font = NOTE

    ws['A8'] = '2. いま、月にどれくらい時間を使っているか'
    ws['A8'].font = H2
    for col, label in [('A', '作業'), ('B', '1案件あたり'), ('C', '全案件'), ('D', '人件費')]:
        c = ws[f'{col}9']
        c.value, c.font, c.fill, c.border = label, H2, GRAY, BOX
    ws['E9'] = '「1案件あたり」の時間を、実感に合わせて書き換えてください'
    ws['E9'].font = NOTE

    r = 10
    for item in tco['hoursPerProject']:
        ws[f'A{r}'] = item['label']
        c = ws[f'B{r}']
        c.value, c.font, c.fill, c.border, c.number_format = item['hours'], BLUE, YELLOW, BOX, HRS
        ws[f'C{r}'] = f'=B{r}*$B$6'
        ws[f'C{r}'].font, ws[f'C{r}'].number_format = BLACK, HRS
        ws[f'D{r}'] = f'=C{r}*$B$5'
        ws[f'D{r}'].font, ws[f'D{r}'].number_format = BLACK, YEN
        ws[f'E{r}'] = item['note']
        ws[f'E{r}'].font = NOTE
        r += 1

    last = r - 1
    ws[f'A{r}'] = '合計'
    ws[f'A{r}'].font = H2
    ws[f'C{r}'] = f'=SUM(C10:C{last})'
    ws[f'C{r}'].font, ws[f'C{r}'].number_format = BOLD, HRS
    ws[f'D{r}'] = f'=SUM(D10:D{last})'
    ws[f'D{r}'].font, ws[f'D{r}'].number_format = BOLD, YEN
    for col in 'ABCD':
        ws[f'{col}{r}'].fill, ws[f'{col}{r}'].border = GRAY, BOX
    total = r

    r += 2
    ws[f'A{r}'] = '3. 導入したあと、どれくらい減るか'
    ws[f'A{r}'].font = H2
    r += 1
    ws[f'A{r}'] = '減らせる割合'
    c = ws[f'B{r}']
    c.value, c.font, c.fill, c.border, c.number_format = tco['reductionRate'], BLUE, YELLOW, BOX, '0%'
    ws[f'E{r}'] = '控えめに見た割合。報告と転記が要らなくなる分です'
    ws[f'E{r}'].font = NOTE
    rate = r

    r += 1
    ws[f'A{r}'] = '減らせる人件費（月）'
    ws[f'D{r}'] = f'=D{total}*B{rate}'
    ws[f'D{r}'].font, ws[f'D{r}'].number_format = BLACK, YEN
    saved = r

    r += 1
    ws[f'A{r}'] = f'AgentPM の月額（{F["taxNote"]}）'
    c = ws[f'D{r}']
    c.value, c.font, c.fill, c.border, c.number_format = pro['monthlyYen'], BLUE, YELLOW, BOX, YEN
    ws[f'E{r}'] = f'{pro["label"]}（{pro["maxMembers"]}名・{pro["maxProjects"]}プロジェクトまで）。無料プランなら0円'
    ws[f'E{r}'].font = NOTE
    fee = r

    r += 2
    ws[f'A{r}'] = '差し引き（月）'
    ws[f'A{r}'].font = H2
    ws[f'D{r}'] = f'=D{saved}-D{fee}'
    ws[f'D{r}'].font = Font(name='Arial', bold=True, size=12)
    ws[f'D{r}'].number_format, ws[f'D{r}'].fill, ws[f'D{r}'].border = YEN, GRAY, BOX
    month = r
    r += 1
    ws[f'A{r}'] = '差し引き（年）'
    ws[f'A{r}'].font = H2
    ws[f'D{r}'] = f'=D{month}*12'
    ws[f'D{r}'].font = Font(name='Arial', bold=True, size=12)
    ws[f'D{r}'].number_format, ws[f'D{r}'].fill, ws[f'D{r}'].border = YEN, GRAY, BOX

    r += 2
    for line in [
        '※ 時間の初期値は、お客様から伺った範囲での目安です。実測値ではありません。',
        f'※ 月額は{F["asOf"].replace("-", "年")}月時点の{F["taxNote"]}です。最新は agentpm.app/pricing をご確認ください。',
        '※ 相手先（クライアント）の人数は、どのプランでも料金に影響しません。',
    ]:
        ws[f'A{r}'] = line
        ws[f'A{r}'].font = NOTE
        r += 1

    ws['B5'].comment = Comment('自社のスタッフの平均時給。相手先の人件費は含めません。', 'AgentPM')
    ws[f'B{rate}'].comment = Comment('報告と転記が要らなくなる分を見込んだ割合です。実績ではありません。', 'AgentPM')
    wb.save(OUT / 'agentpm-roi.xlsx')
    print('  agentpm-roi.xlsx')


def build_migration() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = '移行計画'
    ws.sheet_view.showGridLines = False
    for col, w in zip('ABCDEF', [8, 40, 16, 14, 14, 44]):
        ws.column_dimensions[col].width = w

    ws['A1'] = 'AgentPM 移行計画（雛形）'
    ws['A1'].font = H
    ws['A2'] = '黄色い枠（担当・予定日・状態）を埋めてお使いください。週の区切りは目安です。'
    ws['A2'].font = NOTE
    ws['A3'] = 'いま使っているツールは、移し終えるまで止めなくて構いません。並べて使えます。'
    ws['A3'].font = NOTE

    for col, label in [('A', '週'), ('B', 'やること'), ('C', '担当'), ('D', '予定日'), ('E', '状態'), ('F', '補足')]:
        c = ws[f'{col}5']
        c.value, c.font, c.fill, c.border = label, WHITE_B, NAVY, BOX

    free = F['agentpm']['free']
    tasks = [
        ('1週目', '無料プランで登録する', 'クレジットカードは要りません'),
        ('1週目', '試す案件を1つ決める', '進行中の小さめの案件が向いています'),
        ('1週目', 'テンプレートからプロジェクトを作る', '業種を選ぶとWikiとマイルストーンが入った状態で立ち上がります'),
        ('1週目', 'いまのツールからCSVで書き出す', 'タスクの一覧を出しておきます'),
        ('2週目', 'CSVを取り込む', '取り込む前に、何件をどう登録するかが一覧で出ます'),
        ('2週目', '担当者と期限を入れ直す', '取り込みでは入りきらない項目を補います'),
        ('2週目', '社内メンバーを招待する', f'無料プランは{free["maxMembers"]}名まで。まず3〜5人で'),
        ('2週目', '「いま誰の番か」を全員で1周確認する', 'ここがこのツールの肝です'),
        ('3週目', '相手先を招待する', '招待メールから。人数は何人でも無料'),
        ('3週目', '見せるタスクを決める', 'はじめは見せない設定です'),
        ('3週目', '相手先に使い方を一言伝える', '「招待メールから入れば進捗が見えます」で足ります'),
        ('3週目', '最初の承認を1件回してみる', 'メールのリンクから押せます'),
        ('4週目', 'チャットに秘書を入れる', f'無料プランは{free["maxChatGroups"]}グループまで'),
        ('4週目', '拾い方を決める', '毎時まとめて／メンション時のみ（即時）／取り込まない'),
        ('4週目', '前のツールを止めるか決める', '併用を続けても構いません'),
        ('4週目', '残りの案件を移す', '1件目で慣れてからにします'),
    ]

    dv = DataValidation(type='list', formula1='"これから,着手,完了,見送り"', allow_blank=True)
    ws.add_data_validation(dv)

    r = 6
    for week, task, note in tasks:
        ws[f'A{r}'] = week
        ws[f'A{r}'].font, ws[f'A{r}'].border = BLACK, BOX
        ws[f'B{r}'] = task
        ws[f'B{r}'].font, ws[f'B{r}'].border, ws[f'B{r}'].alignment = BLACK, BOX, WRAP
        for col in 'CDE':
            c = ws[f'{col}{r}']
            c.font, c.fill, c.border = BLUE, YELLOW, BOX
        ws[f'D{r}'].number_format = 'yyyy/mm/dd'
        dv.add(ws[f'E{r}'])
        ws[f'F{r}'] = note
        ws[f'F{r}'].font, ws[f'F{r}'].alignment = NOTE, WRAP
        r += 1

    last = r - 1
    ws[f'A{r + 1}'] = '進み具合'
    ws[f'A{r + 1}'].font = H2
    ws[f'B{r + 1}'] = '完了した数'
    ws[f'C{r + 1}'] = f'=COUNTIF(E6:E{last},"完了")'
    ws[f'C{r + 1}'].font, ws[f'C{r + 1}'].border = BOLD, BOX
    ws[f'B{r + 2}'] = 'ぜんぶで'
    ws[f'C{r + 2}'] = f'=COUNTA(B6:B{last})'
    ws[f'C{r + 2}'].font, ws[f'C{r + 2}'].border = BLACK, BOX
    ws[f'B{r + 3}'] = '達成率'
    ws[f'C{r + 3}'] = f'=IFERROR(C{r + 1}/C{r + 2},0)'
    ws[f'C{r + 3}'].font, ws[f'C{r + 3}'].number_format = BOLD, '0%'
    ws[f'C{r + 3}'].fill, ws[f'C{r + 3}'].border = GRAY, BOX

    ws[f'A{r + 5}'] = '※ この雛形は AgentPM に取り込めます（agentpm task import --file 〜.csv）。'
    ws[f'A{r + 5}'].font = NOTE
    ws[f'A{r + 6}'] = '※ つまずいたら、画面のチャットから聞いてください。移行のお手伝いもします。'
    ws[f'A{r + 6}'].font = NOTE

    wb.save(OUT / 'agentpm-migration-plan.xlsx')
    print('  agentpm-migration-plan.xlsx')


build_roi()
build_migration()
